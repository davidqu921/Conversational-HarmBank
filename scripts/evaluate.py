"""
evaluate.py
===========
Compare LLM coding to the current human truth labels.

The current hand labels live in `data_prep/sample_truth_label.csv`, a wide
consensus sheet with one column per label and `X` marking presence. The loader
also keeps backward-compatible support for the older long CSV format.

Outputs:
  1. agreement_report.md
  2. confusion_matrices.xlsx
  3. discrepancies.xlsx
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
from collections import Counter
from pathlib import Path


ATTACK_VECTORS = ["Brute Force", "Disguised Intent", "Role Play", "Structured Response", "AI Attack", "No Attempt"]
SUBTYPES = [
    "One Shot", "Overwhelming", "Out of Context",
    "Testing", "Joking", "Wordplay",
    "Fake Persona", "Emotional Manipulation", "Threats", "Fake Scenario",
    "Different Language", "Text Continuation",
    "Privilages", "Re-prompting", "Appeal to underlying system", "No Attempt",
]
ATTEMPTS = [
    "Swearing", "Slur", "Inappropriate Topics", "Bias", "Political/Social",
    "Incorrect Information", "Break Character", "Abnormal", "General Break",
]
CONV_LABELS = ["Single Prompt", "Multi Prompt", "Conversational"]
SUCCESS_LABELS = ["Yes - Broken", "Maybe", "No - Safe"]
SEVERITY_LABELS = ["0 - Safe", "1 - Minor", "2 - Major", "3 - Severe"]

SUBTYPE_TO_VECTOR = {
    "One Shot": "Brute Force",
    "Overwhelming": "Brute Force",
    "Out of Context": "Brute Force",
    "Testing": "Disguised Intent",
    "Joking": "Disguised Intent",
    "Wordplay": "Disguised Intent",
    "Fake Persona": "Role Play",
    "Emotional Manipulation": "Role Play",
    "Threats": "Role Play",
    "Fake Scenario": "Role Play",
    "Different Language": "Structured Response",
    "Text Continuation": "Structured Response",
    "Privilages": "AI Attack",
    "Re-prompting": "AI Attack",
    "Appeal to underlying system": "AI Attack",
    "No Attempt": "No Attempt",
}


def normalize(s: str | None) -> str:
    if not s:
        return ""
    s = " ".join(str(s).strip().split())
    aliases = {
        "structured responses": "Structured Response",
        "structured response": "Structured Response",
        "no attempt": "No Attempt",
        "out of context": "Out of Context",
        "diff language": "Different Language",
        "different language": "Different Language",
        "inapp topic": "Inappropriate Topics",
        "inappropriate topics": "Inappropriate Topics",
        "incorrect info": "Incorrect Information",
        "incorrect information": "Incorrect Information",
        "general break": "General Break",
        "appeal to ai": "Appeal to underlying system",
        "appeal to underlying ai": "Appeal to underlying system",
        "appeal to underlying system": "Appeal to underlying system",
        "privileges": "Privilages",
        "privilages": "Privilages",
        "minor (silly unsual)": "1 - Minor",
        "minor (silly unusual)": "1 - Minor",
        "major (problematic)": "2 - Major",
        "definitely broken (severe)": "3 - Severe",
        "no (safe)": "0 - Safe",
        "0 - safe": "0 - Safe",
        "1- minor": "1 - Minor",
        "1 - minor": "1 - Minor",
        "2 - major": "2 - Major",
        "3 - severe": "3 - Severe",
    }
    return aliases.get(s.lower(), s)


def empty_record() -> dict:
    return {"vectors": set(), "subtypes": set(), "attempts": set(), "conv": "", "success": "", "severity": "", "reasoning": ""}


def load_llm(jsonl_path: Path) -> dict[str, dict]:
    out: dict[str, dict] = {}
    with jsonl_path.open(encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            rec = json.loads(line)
            if rec.get("error") or not rec.get("llm_output"):
                continue
            cid = str(rec["id"]).strip()
            o = rec["llm_output"]
            entry = empty_record()
            for coding in o.get("codings", []):
                vec = normalize(coding.get("attack_vector"))
                if vec:
                    entry["vectors"].add(vec)
                for sub in coding.get("subtypes") or []:
                    code = normalize(sub.get("code") if isinstance(sub, dict) else sub)
                    if code and code != "Other":
                        entry["subtypes"].add(code)
                for attempt in coding.get("attempts") or []:
                    attempt = normalize(attempt)
                    if attempt and attempt != "Other":
                        entry["attempts"].add(attempt)
            if "No Attempt" in entry["vectors"] and not entry["subtypes"]:
                entry["subtypes"].add("No Attempt")
            entry["conv"] = normalize(o.get("conversational"))
            entry["success"] = normalize(o.get("success"))
            entry["severity"] = normalize(o.get("severity"))
            if not entry["severity"] and entry["success"] == "No - Safe":
                entry["severity"] = "0 - Safe"
            entry["reasoning"] = o.get("reasoning", "")
            out[cid] = entry
    return out


def load_human(csv_path: Path) -> dict[tuple[str, str], dict]:
    with csv_path.open(encoding="utf-8-sig", newline="") as f:
        sample = list(csv.reader(f))
    if not sample:
        return {}
    if sample[0] and sample[0][0] == "Interaction":
        return load_wide_truth(sample)
    return load_long_truth(csv_path)


def load_wide_truth(rows: list[list[str]]) -> dict[tuple[str, str], dict]:
    labels = rows[2]
    out: dict[tuple[str, str], dict] = {}

    for row in rows[3:]:
        if not row or not row[0].strip():
            continue
        cid = row[0].strip()
        entry = empty_record()
        for i, raw_value in enumerate(row[1:], 1):
            if raw_value.strip().upper() != "X" or i >= len(labels):
                continue
            label = normalize(labels[i])
            if not label:
                label = "No Attempt"
            if i == 1:
                entry["vectors"].add("No Attempt")
                entry["subtypes"].add("No Attempt")
            elif label in SUBTYPES:
                entry["subtypes"].add(label)
                parent = SUBTYPE_TO_VECTOR.get(label)
                if parent:
                    entry["vectors"].add(parent)
            elif label in ATTEMPTS:
                entry["attempts"].add(label)
            elif label in CONV_LABELS:
                entry["conv"] = label
            elif label in SUCCESS_LABELS:
                entry["success"] = label
            elif label in SEVERITY_LABELS:
                entry["severity"] = label
        if entry["success"] == "No - Safe" and not entry["severity"]:
            entry["severity"] = "0 - Safe"
        out[("Consensus", cid)] = entry
    return out


def load_long_truth(csv_path: Path) -> dict[tuple[str, str], dict]:
    out: dict[tuple[str, str], dict] = {}
    with csv_path.open(encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            coder = row.get("coder", "Consensus") or "Consensus"
            cid = row.get("id") or row.get("Interaction")
            if not cid:
                continue
            key = (coder, str(cid).strip())
            entry = out.setdefault(key, empty_record())
            vec = normalize(row.get("attack_vector"))
            subtype = normalize(row.get("subtype"))
            attempt = normalize(row.get("attempt"))
            if vec:
                entry["vectors"].add(vec)
            if subtype:
                entry["subtypes"].add(subtype)
            if attempt and attempt != "No Attempt":
                entry["attempts"].add(attempt)
            entry["conv"] = entry["conv"] or normalize(row.get("conversational"))
            entry["success"] = entry["success"] or normalize(row.get("success"))
            entry["severity"] = entry["severity"] or normalize(row.get("severity"))
            if entry["success"] == "No - Safe" and not entry["severity"]:
                entry["severity"] = "0 - Safe"
    return out


def load_transcripts(jsonl_path: Path) -> dict[str, str]:
    out = {}
    with jsonl_path.open(encoding="utf-8") as f:
        for line in f:
            if line.strip():
                rec = json.loads(line)
                out[str(rec["id"])] = rec.get("transcript_text", "")
    return out


def cohen_kappa(a: list[str], b: list[str]) -> float:
    assert len(a) == len(b) and len(a) > 0
    labels = sorted(set(a) | set(b))
    n = len(a)
    po = sum(1 for x, y in zip(a, b) if x == y) / n
    ca = Counter(a)
    cb = Counter(b)
    pe = sum((ca[label] / n) * (cb[label] / n) for label in labels)
    return 1.0 if pe >= 1.0 else (po - pe) / (1 - pe)


def binary_kappa_per_label(label_pool: list[str], a_sets: list[set], b_sets: list[set]) -> dict[str, float]:
    out = {}
    for label in label_pool:
        a = ["Y" if label in labels else "N" for labels in a_sets]
        b = ["Y" if label in labels else "N" for labels in b_sets]
        out[label] = float("nan") if set(a) == {"N"} and set(b) == {"N"} else cohen_kappa(a, b)
    return out


def macro_kappa(per_label: dict[str, float]) -> float:
    vals = [v for v in per_label.values() if v == v]
    return sum(vals) / len(vals) if vals else float("nan")


def first_label(labels: set, label_pool: list[str]) -> str:
    for label in label_pool:
        if label in labels:
            return label
    return ""


def confusion_matrix(rows_a: list[str], rows_b: list[str], labels: list[str]) -> list[list[int]]:
    idx = {label: i for i, label in enumerate(labels)}
    matrix = [[0] * len(labels) for _ in labels]
    for a, b in zip(rows_a, rows_b):
        if a in idx and b in idx:
            matrix[idx[a]][idx[b]] += 1
    return matrix


def multilabel_horizontal_metrics(human_sets: list[set], llm_sets: list[set]) -> dict[str, float]:
    """Conversation-level set agreement for one multi-label dimension."""
    n = len(human_sets)
    if n == 0:
        return {
            "hit_rate": float("nan"),
            "exact": float("nan"),
            "jaccard": float("nan"),
            "precision": float("nan"),
            "recall": float("nan"),
            "f1": float("nan"),
        }

    hits = []
    exact = []
    jaccards = []
    precisions = []
    recalls = []
    f1s = []

    for human, llm in zip(human_sets, llm_sets):
        intersection = human & llm
        union = human | llm
        hits.append(1.0 if (not human and not llm) or bool(intersection) else 0.0)
        exact.append(1.0 if human == llm else 0.0)
        jaccards.append(1.0 if not union else len(intersection) / len(union))
        precision = 1.0 if not llm else len(intersection) / len(llm)
        recall = 1.0 if not human else len(intersection) / len(human)
        f1 = 0.0 if precision + recall == 0 else 2 * precision * recall / (precision + recall)
        precisions.append(precision)
        recalls.append(recall)
        f1s.append(f1)

    return {
        "hit_rate": sum(hits) / n,
        "exact": sum(exact) / n,
        "jaccard": sum(jaccards) / n,
        "precision": sum(precisions) / n,
        "recall": sum(recalls) / n,
        "f1": sum(f1s) / n,
    }


def single_label_hit_rate(human_values: list[str], llm_values: list[str]) -> float:
    if not human_values:
        return float("nan")
    return sum(1 for human, llm in zip(human_values, llm_values) if human == llm) / len(human_values)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--llm-codings", required=True, type=Path)
    parser.add_argument("--human-coding", required=True, type=Path)
    parser.add_argument("--transcripts", required=True, type=Path)
    parser.add_argument("--out-dir", required=True, type=Path)
    args = parser.parse_args()

    args.out_dir.mkdir(parents=True, exist_ok=True)
    llm = load_llm(args.llm_codings)
    human = load_human(args.human_coding)
    transcripts = load_transcripts(args.transcripts)

    coders = sorted({coder for coder, _ in human})
    coded_ids = sorted({cid for _, cid in human})
    eval_ids = sorted(set(coded_ids) & set(llm))
    if not eval_ids:
        sys.exit("no overlap between LLM coding and human coding - cannot evaluate")

    report_lines = [
        "# Agreement Report",
        "",
        f"- Eval set: **{len(eval_ids)} conversations** with both LLM and human coding",
        f"- Human label source: `{args.human_coding}`",
        f"- Human coders/label sets: {', '.join(coders)}",
        "",
        "## Cohen's kappa per dimension",
        "",
        "| Coder | n | Vector (macro) | Type (macro) | Attempt (macro) | Conversational | Success | Severity |",
        "|---|---:|---:|---:|---:|---:|---:|---:|",
    ]

    h_all, l_all = [], []
    for coder in coders:
        h, l = [], []
        for cid in eval_ids:
            key = (coder, cid)
            if key in human:
                h.append(human[key])
                l.append(llm[cid])
        if not h:
            continue
        h_all.extend(h)
        l_all.extend(l)
        kv = macro_kappa(binary_kappa_per_label(ATTACK_VECTORS, [r["vectors"] for r in h], [r["vectors"] for r in l]))
        kt = macro_kappa(binary_kappa_per_label(SUBTYPES, [r["subtypes"] for r in h], [r["subtypes"] for r in l]))
        ka = macro_kappa(binary_kappa_per_label(ATTEMPTS, [r["attempts"] for r in h], [r["attempts"] for r in l]))
        kc = cohen_kappa([r["conv"] or "<blank>" for r in h], [r["conv"] or "<blank>" for r in l])
        ks = cohen_kappa([r["success"] or "<blank>" for r in h], [r["success"] or "<blank>" for r in l])
        kz = cohen_kappa([r["severity"] or "<blank>" for r in h], [r["severity"] or "<blank>" for r in l])
        report_lines.append(f"| {coder} | {len(h)} | {kv:.3f} | {kt:.3f} | {ka:.3f} | {kc:.3f} | {ks:.3f} | {kz:.3f} |")

    kv = macro_kappa(binary_kappa_per_label(ATTACK_VECTORS, [r["vectors"] for r in h_all], [r["vectors"] for r in l_all]))
    kt = macro_kappa(binary_kappa_per_label(SUBTYPES, [r["subtypes"] for r in h_all], [r["subtypes"] for r in l_all]))
    ka = macro_kappa(binary_kappa_per_label(ATTEMPTS, [r["attempts"] for r in h_all], [r["attempts"] for r in l_all]))
    kc = cohen_kappa([r["conv"] or "<blank>" for r in h_all], [r["conv"] or "<blank>" for r in l_all])
    ks = cohen_kappa([r["success"] or "<blank>" for r in h_all], [r["success"] or "<blank>" for r in l_all])
    kz = cohen_kappa([r["severity"] or "<blank>" for r in h_all], [r["severity"] or "<blank>" for r in l_all])
    report_lines.append(f"| **Pooled** | **{len(h_all)}** | **{kv:.3f}** | **{kt:.3f}** | **{ka:.3f}** | **{kc:.3f}** | **{ks:.3f}** | **{kz:.3f}** |")

    report_lines.extend(["", "## Per-label kappa", "", "### Attack Vector", "| Label | kappa | LLM positives | Human positives |", "|---|---:|---:|---:|"])
    for label, k in binary_kappa_per_label(ATTACK_VECTORS, [r["vectors"] for r in h_all], [r["vectors"] for r in l_all]).items():
        report_lines.append(f"| {label} | {k:.3f} | {sum(1 for r in l_all if label in r['vectors'])} | {sum(1 for r in h_all if label in r['vectors'])} |")

    report_lines.extend(["", "### Severity", "| Label | LLM positives | Human positives |", "|---|---:|---:|"])
    for label in SEVERITY_LABELS:
        report_lines.append(f"| {label} | {sum(1 for r in l_all if r['severity'] == label)} | {sum(1 for r in h_all if r['severity'] == label)} |")

    report_lines.extend([
        "",
        "## Horizontal conversation-level agreement",
        "",
        "These metrics compare each conversation's whole label set within one dimension. They are complementary to the label-wise Cohen's kappa above.",
        "",
        "### Multi-label dimensions",
        "",
        "| Dimension | Hit Rate | Exact Match | Mean Jaccard | Mean Precision | Mean Recall | Mean F1 |",
        "|---|---:|---:|---:|---:|---:|---:|",
    ])
    for name, field in [("Vector", "vectors"), ("Type", "subtypes"), ("Attempt", "attempts")]:
        metrics = multilabel_horizontal_metrics([r[field] for r in h_all], [r[field] for r in l_all])
        report_lines.append(
            f"| {name} | {metrics['hit_rate']:.3f} | {metrics['exact']:.3f} | {metrics['jaccard']:.3f} | "
            f"{metrics['precision']:.3f} | {metrics['recall']:.3f} | {metrics['f1']:.3f} |"
        )

    report_lines.extend([
        "",
        "### Single-label dimensions",
        "",
        "| Dimension | Hit Rate |",
        "|---|---:|",
    ])
    for name, field in [("Conversational", "conv"), ("Success", "success"), ("Severity", "severity")]:
        hit_rate = single_label_hit_rate([r[field] or "<blank>" for r in h_all], [r[field] or "<blank>" for r in l_all])
        report_lines.append(f"| {name} | {hit_rate:.3f} |")

    (args.out_dir / "agreement_report.md").write_text("\n".join(report_lines), encoding="utf-8")
    print(f"  -> {args.out_dir / 'agreement_report.md'}", file=sys.stderr)

    try:
        import openpyxl

        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        def write_matrix(sheet_name: str, labels: list[str], llm_rows: list[str], human_rows: list[str]):
            ws = wb.create_sheet(sheet_name)
            ws.cell(1, 1, f"{sheet_name}: rows=LLM, cols=human")
            for j, label in enumerate(labels, 2):
                ws.cell(2, j, label)
            for i, label in enumerate(labels, 3):
                ws.cell(i, 1, label)
            for i, row in enumerate(confusion_matrix(llm_rows, human_rows, labels), 3):
                for j, value in enumerate(row, 2):
                    ws.cell(i, j, value)

        write_matrix("vector", ATTACK_VECTORS, [first_label(r["vectors"], ATTACK_VECTORS) for r in l_all], [first_label(r["vectors"], ATTACK_VECTORS) for r in h_all])
        write_matrix("type", SUBTYPES, [first_label(r["subtypes"], SUBTYPES) for r in l_all], [first_label(r["subtypes"], SUBTYPES) for r in h_all])
        write_matrix("attempt", ATTEMPTS, [first_label(r["attempts"], ATTEMPTS) for r in l_all], [first_label(r["attempts"], ATTEMPTS) for r in h_all])
        write_matrix("conversational", CONV_LABELS, [r["conv"] for r in l_all], [r["conv"] for r in h_all])
        write_matrix("success", SUCCESS_LABELS, [r["success"] for r in l_all], [r["success"] for r in h_all])
        write_matrix("severity", SEVERITY_LABELS, [r["severity"] for r in l_all], [r["severity"] for r in h_all])
        wb.save(args.out_dir / "confusion_matrices.xlsx")
        print(f"  -> {args.out_dir / 'confusion_matrices.xlsx'}", file=sys.stderr)
    except ImportError:
        print("  (skipping confusion matrices - install openpyxl)", file=sys.stderr)

    try:
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Discrepancies"
        headers = [
            "conversation_id", "coder", "human_vectors", "llm_vectors", "human_types", "llm_types",
            "human_attempts", "llm_attempts", "human_conv", "llm_conv", "human_success", "llm_success",
            "human_severity", "llm_severity", "agree_vector", "agree_type", "agree_success",
            "agree_severity", "llm_reasoning", "transcript",
        ]
        for col, header in enumerate(headers, 1):
            ws.cell(1, col, header)
        row_num = 2
        for coder in coders:
            for cid in eval_ids:
                key = (coder, cid)
                if key not in human:
                    continue
                h = human[key]
                l = llm[cid]
                agree_vector = bool(h["vectors"] & l["vectors"])
                agree_type = bool(h["subtypes"] & l["subtypes"]) or (not h["subtypes"] and not l["subtypes"])
                agree_success = h["success"] == l["success"]
                agree_severity = h["severity"] == l["severity"]
                if agree_vector and agree_type and agree_success and agree_severity:
                    continue
                values = [
                    cid, coder,
                    "; ".join(sorted(h["vectors"])), "; ".join(sorted(l["vectors"])),
                    "; ".join(sorted(h["subtypes"])), "; ".join(sorted(l["subtypes"])),
                    "; ".join(sorted(h["attempts"])), "; ".join(sorted(l["attempts"])),
                    h["conv"], l["conv"], h["success"], l["success"], h["severity"], l["severity"],
                    agree_vector, agree_type, agree_success, agree_severity,
                    l["reasoning"][:600], transcripts.get(cid, "")[:3000],
                ]
                for col, value in enumerate(values, 1):
                    ws.cell(row_num, col, value)
                row_num += 1
        wb.save(args.out_dir / "discrepancies.xlsx")
        print(f"  -> {args.out_dir / 'discrepancies.xlsx'}  ({row_num - 2} disagreements)", file=sys.stderr)
    except ImportError:
        print("  (skipping discrepancies xlsx - install openpyxl)", file=sys.stderr)


if __name__ == "__main__":
    main()

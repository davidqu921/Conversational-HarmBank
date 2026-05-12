"""
prepare_data.py
================
Loads raw transcripts and human-coded ground truth into normalized formats the
other scripts in this skill consume.

Outputs (written under --out-dir, default ./data_prep/):
- transcripts.jsonl   one line per conversation: {id, created_time, report, transcript_text, n_turns}
- ground_truth.csv    long-format human codings (one row per (coder, conversation, attack_vector_index)
                      from the per-coder sheets that use the *current* codebook)
- irr_set.json        list of conversation IDs coded by >=3 of the current-codebook coders
                      (these are the IRR conversations we use to compute human-vs-human and
                      human-vs-LLM agreement)

CLI:
    python -m scripts.prepare_data \\
        --transcripts-csv /path/to/breakme-steve_transcripts.csv \\
        --coding-xlsx     /path/to/VirtualSteve\\ Coding\\ -\\ Round\\ 2.xlsx \\
        --out-dir         ./data_prep
"""
from __future__ import annotations
import argparse
import csv
import json
import sys
from collections import defaultdict
from pathlib import Path

# Sheets in the human-coding XLSX that use the *current* codebook (matches references/codebook.md).
# Older sheets ("Coding First Attempt", "Coding Round 2", etc.) use a different/older taxonomy
# (e.g. "Simulation", "AI Privelages", "Hybrid") and are deliberately excluded.
CURRENT_CODEBOOK_SHEETS = [
    "Coding - Jailbreak - David",
    "Coding - Jailbreak - Taylor",
    "Coding - Jailbreak-Anushree",
    "Coding - Jailbreak - Anson Kwok",
    "Coding - Jailbreak - Jania",
]

# Map sheet name -> short coder name
def coder_from_sheet(sheet_name: str) -> str:
    # 'Coding - Jailbreak - David' -> 'David'
    # 'Coding - Jailbreak-Anushree' -> 'Anushree' (note hyphen, not " - ")
    name = sheet_name.split(" - ")[-1]
    if "-" in name and not name.startswith("Anson"):  # 'Jailbreak-Anushree' case
        name = name.split("-")[-1]
    return name.strip()


def format_transcript(chat_json_str: str) -> tuple[str, int]:
    """Convert the raw chatTranscript JSON (alternating role/content turns) into a
    human-readable transcript and a turn count.

    The raw format has many sub-turn fragments (voice-to-text artifacts) where
    the same speaker contributes multiple consecutive entries; we coalesce these
    into one line per speaker-shift to make the transcript easier for the LLM to read.
    """
    turns = json.loads(chat_json_str)
    lines: list[str] = []
    cur_role = None
    cur_buf: list[str] = []
    n_speaker_shifts = 0

    def flush():
        nonlocal cur_buf
        if cur_role is not None and cur_buf:
            speaker = "Steve" if cur_role == "assistant" else "Student"
            text = " ".join(s.strip() for s in cur_buf if s.strip())
            if text:
                lines.append(f"{speaker}: {text}")
        cur_buf = []

    for t in turns:
        role = t.get("role")
        content = (t.get("content") or "").strip()
        if not content:
            continue
        if role != cur_role:
            flush()
            cur_role = role
            n_speaker_shifts += 1
        cur_buf.append(content)
    flush()
    return "\n".join(lines), n_speaker_shifts


def load_transcripts(csv_path: Path) -> list[dict]:
    out = []
    with csv_path.open(newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            try:
                transcript_text, n_turns = format_transcript(row["chatTranscript"])
            except Exception as e:
                print(f"  WARN: could not parse transcript for id={row.get('id')}: {e}", file=sys.stderr)
                continue
            report = row.get("report") or ""
            if report.upper() == "NULL":
                report = ""
            out.append({
                "id": str(row["id"]).strip(),
                "created_time": row.get("createdTime", ""),
                "report": report,
                "transcript_text": transcript_text,
                "n_turns": n_turns,
            })
    return out


def load_human_coding(xlsx_path: Path) -> list[dict]:
    """Returns list of {coder, id, attack_vector, subtype, attempt, conversational, success, notes}."""
    import openpyxl  # local import keeps the script importable when openpyxl isn't installed
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    out: list[dict] = []
    for sheet_name in CURRENT_CODEBOOK_SHEETS:
        if sheet_name not in wb.sheetnames:
            print(f"  WARN: sheet '{sheet_name}' not found", file=sys.stderr)
            continue
        ws = wb[sheet_name]
        coder = coder_from_sheet(sheet_name)
        # Detect column count - some sheets have a 'Notes' column at index 7
        for row in ws.iter_rows(min_row=2, values_only=True):
            cid = row[1]
            vec = row[2]
            if not cid or not vec:
                continue
            cid_str = str(cid).replace(".0", "").strip()
            out.append({
                "coder": coder,
                "id": cid_str,
                "attack_vector": (str(vec) if vec is not None else "").strip(),
                "subtype": (str(row[3]) if row[3] is not None else "").strip(),
                "attempt": (str(row[4]) if row[4] is not None else "").strip(),
                "conversational": (str(row[5]) if row[5] is not None else "").strip(),
                "success": (str(row[6]) if row[6] is not None else "").strip(),
                "notes": (str(row[7]) if len(row) > 7 and row[7] is not None else "").strip(),
            })
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--transcripts-csv", required=True, type=Path)
    ap.add_argument("--coding-xlsx", required=True, type=Path)
    ap.add_argument("--out-dir", default=Path("data_prep"), type=Path)
    args = ap.parse_args()

    args.out_dir.mkdir(parents=True, exist_ok=True)

    print(f"Loading transcripts from {args.transcripts_csv} ...")
    transcripts = load_transcripts(args.transcripts_csv)
    print(f"  {len(transcripts)} transcripts loaded")

    out_jsonl = args.out_dir / "transcripts.jsonl"
    with out_jsonl.open("w", encoding="utf-8") as f:
        for t in transcripts:
            f.write(json.dumps(t) + "\n")
    print(f"  -> {out_jsonl}")

    print(f"\nLoading human coding from {args.coding_xlsx} ...")
    coding = load_human_coding(args.coding_xlsx)
    print(f"  {len(coding)} rows from {len({c['coder'] for c in coding})} coders covering {len({c['id'] for c in coding})} unique conversations")

    out_csv = args.out_dir / "ground_truth.csv"
    with out_csv.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=["coder", "id", "attack_vector", "subtype", "attempt", "conversational", "success", "notes"])
        writer.writeheader()
        writer.writerows(coding)
    print(f"  -> {out_csv}")

    # IRR set: conversations coded by >=3 of the current-codebook coders
    by_id: dict[str, set] = defaultdict(set)
    for c in coding:
        by_id[c["id"]].add(c["coder"])
    irr_ids = sorted([cid for cid, coders in by_id.items() if len(coders) >= 3])
    out_irr = args.out_dir / "irr_set.json"
    with out_irr.open("w") as f:
        json.dump({"irr_conversation_ids": irr_ids, "n": len(irr_ids)}, f, indent=2)
    print(f"  -> {out_irr} ({len(irr_ids)} IRR conversations)")

    # Summary
    transcript_ids = {t["id"] for t in transcripts}
    coded_ids = {c["id"] for c in coding}
    print("\n=== Summary ===")
    print(f"  Transcripts:           {len(transcripts)}")
    print(f"  Human-coded conversations: {len(coded_ids)}")
    print(f"  Coded conversations IN transcripts:    {len(coded_ids & transcript_ids)}")
    print(f"  Coded conversations NOT in transcripts: {len(coded_ids - transcript_ids)}")
    print(f"  IRR set (>=3 coders):  {len(irr_ids)}")


if __name__ == "__main__":
    main()
